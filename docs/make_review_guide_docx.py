# -*- coding: utf-8 -*-
"""生成 meArm 固件「评审向说明文档」docs/meArm_review_guide.docx。

用法:
    python docs/make_review_guide_docx.py [输出路径]

文档内容全部写在同目录的 review_content.py 里（DOC 列表），本文件只负责渲染:
    DOC = [ ("h1"|"h2"|"h3", 文本),
            ("p", 文本),                       # 支持 **粗体** 标记
            ("bul"|"num", [文本, ...]),
            ("tbl", 表题, [表头...], [[单元格...], ...], [列宽cm...]),
            ("pb",)                            # 分页
          ]
"""
import os
import re
import sys
import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import review_content  # noqa: E402

BODY_FONT = "宋体"
HEAD_FONT = "黑体"


def set_run_font(run, font=BODY_FONT, size=10.5, bold=False, italic=False,
                 color=None):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    rfonts.set(qn("w:eastAsia"), font)


def add_rich(par, text, font=BODY_FONT, size=10.5, base_bold=False):
    """把 **粗体** 与 `等宽` 标记渲染成真正的 run。"""
    for token in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if token == "":
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            # 允许 **`等宽`** 这种嵌套：交给同函数递归处理内层标记
            add_rich(par, token[2:-2], font=font, size=size, base_bold=True)
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = par.add_run(token[1:-1])
            set_run_font(run, font="Consolas", size=size - 0.5,
                         bold=base_bold)
        else:
            run = par.add_run(token)
            set_run_font(run, font=font, size=size, bold=base_bold)


def add_par(doc, text, size=10.5, align=None, space_after=4, indent=0.0,
            line=15.5):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing = Pt(line)
    if indent:
        pf.left_indent = Cm(indent)
    if align is not None:
        par.alignment = align
    add_rich(par, text, size=size)
    return par


def add_heading(doc, level, text):
    sizes = {1: 16, 2: 13.5, 3: 12}
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_before = Pt(12 if level == 1 else 8)
    pf.space_after = Pt(5)
    pf.line_spacing = Pt(19 if level == 1 else 17)
    if level == 1:
        pf.page_break_before = False
    add_rich(par, text, font=HEAD_FONT, size=sizes[level], base_bold=True)
    return par


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def add_table(doc, caption, headers, rows, widths=None):
    if caption:
        add_par(doc, caption, size=10, align=WD_ALIGN_PARAGRAPH.LEFT,
                space_after=2)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        par = hdr[i].paragraphs[0]
        par.paragraph_format.space_after = Pt(1)
        par.paragraph_format.line_spacing = Pt(13)
        add_rich(par, h, size=9.5, base_bold=True)
        shade_cell(hdr[i], "E8E8E8")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.line_spacing = Pt(13)
            add_rich(par, str(val), size=9.5)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)
    add_par(doc, "", size=6, space_after=2)
    return table


def build(out_path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    for item in review_content.DOC:
        tag = item[0]
        if tag == "h1":
            add_heading(doc, 1, item[1])
        elif tag == "h2":
            add_heading(doc, 2, item[1])
        elif tag == "h3":
            add_heading(doc, 3, item[1])
        elif tag == "p":
            add_par(doc, item[1])
        elif tag == "pc":
            add_par(doc, item[1], align=WD_ALIGN_PARAGRAPH.CENTER)
        elif tag == "code":
            par = add_par(doc, item[1], size=9.5, indent=0.6, space_after=4)
            for run in par.runs:
                set_run_font(run, font="Consolas", size=9.5)
        elif tag == "bul":
            for t in item[1]:
                add_par(doc, "· " + t, indent=0.5, space_after=2)
        elif tag == "num":
            for i, t in enumerate(item[1], 1):
                add_par(doc, "%d) %s" % (i, t), indent=0.5, space_after=2)
        elif tag == "tbl":
            add_table(doc, item[1], item[2], item[3],
                      item[4] if len(item) > 4 else None)
        elif tag == "title":
            add_par(doc, item[1], size=22, align=WD_ALIGN_PARAGRAPH.CENTER,
                    space_after=6, line=30)
        elif tag == "pb":
            par = doc.add_paragraph()
            par.add_run().add_break(WD_BREAK.PAGE)
        else:
            raise ValueError("unknown tag: %r" % (tag,))

    doc.save(out_path)
    return out_path


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        HERE, "meArm_review_guide.docx")
    build(out)
    size = os.path.getsize(out)
    print("docx written: %s (%d bytes) %s" % (out, size,
                                              datetime.date.today().isoformat()))


if __name__ == "__main__":
    main()
